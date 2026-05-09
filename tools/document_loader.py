from pathlib import Path
from datetime import datetime, timezone

import docx


def load_docx(file_path: str) -> dict:
    doc = docx.Document(file_path)

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)

    tables = []
    for table in doc.tables:
        rows = []
        seen = set()
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            key = tuple(cells)
            if key not in seen:
                seen.add(key)
                rows.append(cells)
        tables.append(rows)

    return {
        "source_file": Path(file_path).name,
        "parsed_at": datetime.now(timezone.utc).isoformat(),
        "text": full_text,
        "paragraphs": paragraphs,
        "paragraph_count": len(paragraphs),
        "char_count": len(full_text),
        "tables": tables,
        "table_count": len(tables),
    }
