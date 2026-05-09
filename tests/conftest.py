import io
import json
import os
import sys
import uuid
from datetime import datetime, timezone
from unittest.mock import MagicMock, patch

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

# ── Stub out strands + mcp before any agent module is imported ──
_mock_strands = MagicMock()
_mock_strands.Agent = MagicMock
_mock_mcp = MagicMock()

sys.modules.setdefault("strands", _mock_strands)
sys.modules.setdefault("strands.models", MagicMock())
sys.modules.setdefault("strands.models.litellm", MagicMock())
sys.modules.setdefault("strands.tools", MagicMock())
sys.modules.setdefault("strands.tools.mcp", MagicMock())
sys.modules.setdefault("mcp", _mock_mcp)
sys.modules.setdefault("mcp.client", MagicMock())
sys.modules.setdefault("mcp.client.stdio", MagicMock())

from fastapi.testclient import TestClient


class FakeDynamoTable:
    """In-memory DynamoDB table mock."""

    def __init__(self):
        self._items: dict = {}

    def put_item(self, Item: dict, **kw):
        self._items[Item["id"]] = dict(Item)

    def get_item(self, Key: dict, **kw):
        item = self._items.get(Key["id"])
        return {"Item": dict(item)} if item else {}

    def delete_item(self, Key: dict, **kw):
        self._items.pop(Key["id"], None)

    def scan(self, **kw):
        items = list(self._items.values())
        limit = kw.get("Limit")
        if limit:
            items = items[:limit]
        return {"Items": items}

    def query(self, **kw):
        items = list(self._items.values())
        expr_values = kw.get("ExpressionAttributeValues", {})
        cid = expr_values.get(":cid")
        if cid:
            items = [i for i in items if i.get("contract_id") == cid]
        limit = kw.get("Limit")
        if limit:
            items = items[:limit]
        return {"Items": items}

    def update_item(self, Key: dict, **kw):
        item = self._items.get(Key["id"], {"id": Key["id"]})
        expr_values = kw.get("ExpressionAttributeValues", {})
        names = kw.get("ExpressionAttributeNames", {})
        update_expr = kw.get("UpdateExpression", "")
        set_part = update_expr.replace("SET ", "", 1) if update_expr.startswith("SET ") else update_expr
        assignments = [a.strip() for a in set_part.split(",")]
        for assignment in assignments:
            if "=" not in assignment:
                continue
            lhs, rhs = [s.strip() for s in assignment.split("=", 1)]
            real_attr = names.get(lhs, lhs)
            value = expr_values.get(rhs)
            if value is not None:
                item[real_attr] = value
        self._items[Key["id"]] = item


@pytest.fixture()
def fake_tables():
    return {
        "contracts": FakeDynamoTable(),
        "reports": FakeDynamoTable(),
        "workflow": FakeDynamoTable(),
        "comments": FakeDynamoTable(),
        "prompts": FakeDynamoTable(),
    }


@pytest.fixture()
def client(fake_tables):
    import api.app as app_module

    app_module._contracts_table = fake_tables["contracts"]
    app_module._reports_table = fake_tables["reports"]
    app_module._workflow_table = fake_tables["workflow"]
    app_module._comments_table = fake_tables["comments"]
    app_module._prompts_table = fake_tables["prompts"]
    app_module._bucket = ""

    return TestClient(app_module.app)


@pytest.fixture()
def sample_docx_bytes():
    from docx import Document

    doc = Document()
    doc.add_heading("소프트웨어 개발 용역 계약서", level=1)
    doc.add_paragraph("제1조 (목적) 본 계약은 메가존클라우드(이하 갑)와 비즈솔루션코리아(이하 을) 간의 소프트웨어 개발 용역에 관한 사항을 정한다.")
    doc.add_paragraph("제2조 (계약기간) 2026년 6월 1일부터 2027년 5월 31일까지로 한다.")
    doc.add_paragraph("제3조 (계약금액) 총 계약금액은 금 5억원(부가세 별도)으로 한다.")
    doc.add_paragraph("제4조 (손해배상) 을은 본 계약 불이행으로 갑에게 손해를 끼친 경우 무제한 배상한다.")
    doc.add_paragraph("제5조 (지식재산권) 본 계약으로 발생하는 모든 지식재산권은 갑에게 완전 이전한다.")
    doc.add_paragraph("제6조 (비밀유지) 양 당사자는 계약 종료 후에도 비밀유지 의무를 진다.")
    doc.add_paragraph("제7조 (해지) 갑은 30일 전 서면 통보로 본 계약을 해지할 수 있다.")
    doc.add_paragraph("제8조 (지체상금) 을이 납기를 지체한 경우 지체일수 1일당 계약금액의 0.1%를 지체상금으로 납부한다.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _make_mock_agent_callable(response_json: str):
    """Create a callable that behaves like a Strands Agent __call__."""
    def agent_call(input_text):
        result = MagicMock()
        result.__str__ = lambda self: response_json
        return result
    return agent_call


@pytest.fixture()
def mock_agents():
    parsing_response = json.dumps({
        "contract_type": "SI",
        "is_standard": False,
        "parties": {
            "party_a": {"name": "메가존클라우드(주)", "representative": None},
            "party_b": {"name": "비즈솔루션코리아", "representative": None},
        },
        "dates": {
            "contract_date": "2026-05-01",
            "start_date": "2026-06-01",
            "end_date": "2027-05-31",
            "renewal_terms": None,
        },
        "financials": {"total_amount": 500000000, "currency": "KRW"},
        "clauses": [],
    }, ensure_ascii=False)

    risk_response = json.dumps({
        "overall_risk": "HIGH",
        "risk_summary": "무제한 배상책임 및 IP 완전이전 조항으로 인해 고위험",
        "clause_risks": [
            {
                "clause_id": "clause_001",
                "risk_level": "HIGH",
                "risk_type": "무제한_배상책임",
                "reason": "배상한도 미설정",
                "recommendation": "배상한도 설정 필요",
                "financial_impact": "최대 5억원",
            },
        ],
        "key_concerns": ["무제한 배상 조항", "IP 완전이전"],
        "escalation_required": True,
    }, ensure_ascii=False)

    legal_response = json.dumps({
        "recommendation": "수정 요청",
        "legal_opinion": "무제한 배상 조항과 IP 완전이전 조항의 수정이 필요합니다.",
        "negotiation_points": ["배상한도를 계약금액의 100%로 제한"],
        "escalation_required": True,
    }, ensure_ascii=False)

    mock_parsing = MagicMock(side_effect=_make_mock_agent_callable(parsing_response))
    mock_risk = MagicMock(side_effect=_make_mock_agent_callable(risk_response))
    mock_legal = MagicMock(side_effect=_make_mock_agent_callable(legal_response))

    import agents.parsing_agent as pa_mod
    import agents.risk_agent as ra_mod
    import agents.legal_review_agent as la_mod

    orig_pa = getattr(pa_mod, "parsing_agent", None)
    orig_ra = getattr(ra_mod, "risk_agent", None)
    orig_la = getattr(la_mod, "legal_review_agent", None)

    pa_mod.parsing_agent = mock_parsing
    ra_mod.risk_agent = mock_risk
    la_mod.legal_review_agent = mock_legal

    yield {"parsing": mock_parsing, "risk": mock_risk, "legal": mock_legal}

    pa_mod.parsing_agent = orig_pa
    ra_mod.risk_agent = orig_ra
    la_mod.legal_review_agent = orig_la
