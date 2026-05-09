import re
from docx import Document


def parse_docx(file_path: str) -> dict:
    """DOCX → Contract JSON (python-docx + regex, LLM 없이)"""
    doc = Document(file_path)
    full_text = "\n".join([p.text for p in doc.paragraphs])

    return {
        "contract_type": detect_type(full_text),
        "is_standard": detect_standard(full_text),
        "parties": extract_parties(full_text),
        "dates": extract_dates(full_text),
        "financials": extract_financials(full_text),
        "clauses": extract_clauses(doc),
    }


def detect_type(text: str) -> str:
    patterns = [
        ("SI",          r"SI도급|시스템구축|소프트웨어개발|개발용역|구축계약|정보시스템"),
        ("Maintenance", r"유지보수계약|하자보수계약|유지관리|기술지원계약"),
        ("SLA",         r"SLA|서비스수준계약|서비스레벨협약|운영서비스계약"),
        ("Outsourcing", r"업무위탁계약|파견계약|인력공급계약|아웃소싱|외주용역"),
        ("MSA",         r"기본계약서|마스터서비스계약|MSA|기본서비스이용계약"),
        ("NDA",         r"비밀유지계약서|비밀유지협약서|NDA|기밀유지계약"),
    ]
    for ctype, pattern in patterns:
        if re.search(pattern, text, re.IGNORECASE):
            return ctype
    return "Other"


def detect_standard(text: str) -> bool:
    return bool(re.search(r"MZC-STD|표준계약서|메가존 표준|MZC 표준|표준 양식", text))


def extract_parties(text: str) -> dict:
    def find_name(role_pattern):
        m = re.search(role_pattern + r"[:\s]*([^\n,()（）]+?)(?:\s*[/,（]|\s*대표)", text)
        return m.group(1).strip() if m else None

    def find_rep(role_pattern):
        m = re.search(role_pattern + r".*?대표이사?\s*([가-힣]{2,5})", text, re.DOTALL)
        return m.group(1).strip() if m else None

    return {
        "party_a": {"name": find_name(r"갑"), "representative": find_rep(r"갑")},
        "party_b": {"name": find_name(r"을"), "representative": find_rep(r"을")},
    }


def extract_dates(text: str) -> dict:
    def to_iso(m):
        if not m:
            return None
        y, mo, d = m.group(1), m.group(2), m.group(3)
        if len(y) == 2:
            y = "20" + y
        return f"{y}-{mo.zfill(2)}-{d.zfill(2)}"

    date_pat = r"(\d{2,4})[년./](\d{1,2})[월./](\d{1,2})일?"

    contract_date = to_iso(re.search(r"(?:계약일|체결일)[^\d]*" + date_pat, text))
    start_date    = to_iso(re.search(r"(?:착수일|시작일|개시일|부터)[^\d]*" + date_pat, text))
    end_date      = to_iso(re.search(r"(?:종료일|만료일|까지)[^\d]*" + date_pat, text))

    renewal_m = re.search(r"(자동갱신|자동연장|묵시적 갱신)[^\n]*\n?[^\n]*", text)
    renewal   = renewal_m.group(0).strip() if renewal_m else None

    return {
        "contract_date": contract_date,
        "start_date": start_date,
        "end_date": end_date,
        "renewal_terms": renewal,
    }


def extract_financials(text: str) -> dict:
    amount = None
    amount_m = re.search(r"[금일]?\s*([\d,]+)\s*원", text)
    if amount_m:
        amount = int(amount_m.group(1).replace(",", ""))

    currency = "KRW"
    if re.search(r"\$|USD", text):
        currency = "USD"
    elif re.search(r"¥|JPY", text):
        currency = "JPY"

    rate = None
    rate_m = re.search(r"([\d.]+)\s*%\s*/?\s*일", text)
    if rate_m:
        rate = float(rate_m.group(1))
    else:
        frac_m = re.search(r"(\d+)\s*/\s*(\d+)", text)
        if frac_m:
            rate = round(int(frac_m.group(1)) / int(frac_m.group(2)) * 100, 4)

    penalty_m = re.search(r"(배상|손해배상|위약금)[^\n]*\n?[^\n]*", text)
    payment_m = re.search(r"(지급|결제|대금)[^\n]*\n?[^\n]*", text)

    return {
        "total_amount": amount,
        "currency": currency,
        "payment_terms": payment_m.group(0).strip() if payment_m else None,
        "penalty_clause": penalty_m.group(0).strip() if penalty_m else None,
        "delay_penalty_rate": rate,
    }


def extract_clauses(doc) -> list:
    """제N조 패턴으로 조항 분리"""
    clauses, current = [], None
    clause_pattern = re.compile(r"^제\s*(\d+)\s*조")

    type_map = {
        r"배상|손해배상|면책|배상한도": "liability",
        r"지식재산권|저작권|특허권|귀속|IP|산출물": "ip",
        r"비밀|기밀|비밀유지|정보보호": "confidentiality",
        r"해지|해제|계약종료|중도해지": "termination",
        r"분쟁|관할|중재|소송|준거법": "dispute",
        r"지체상금|위약금|지연배상|패널티": "penalty",
        r"하자보수|하자담보|무상수리|품질보증": "warranty",
        r"변경요청|범위변경|CR|추가작업|추가개발": "change_request",
        r"자동갱신|자동연장|계약연장|갱신조건": "renewal",
    }

    def classify(title: str) -> str:
        for pattern, ctype in type_map.items():
            if re.search(pattern, title):
                return ctype
        return "other"

    for para in doc.paragraphs:
        match = clause_pattern.match(para.text.strip())
        if match:
            if current:
                clauses.append(current)
            current = {
                "id": f"clause_{match.group(1).zfill(3)}",
                "type": classify(para.text),
                "title": para.text.strip(),
                "content": "",
                "paragraph": int(match.group(1)),
            }
        elif current:
            current["content"] += para.text + "\n"

    if current:
        clauses.append(current)

    for c in clauses:
        c["content"] = c["content"].strip()

    return clauses
